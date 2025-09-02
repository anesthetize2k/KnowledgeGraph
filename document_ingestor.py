import os
import hashlib
from dotenv import load_dotenv
from langchain_community.document_loaders import PyMuPDFLoader, TextLoader, Docx2txtLoader
from langchain.text_splitter import RecursiveCharacterTextSplitter
from neo4j import GraphDatabase
from litellm_wrapper import LiteLLMEmbeddings
from langchain.schema import Document

load_dotenv()


class DocumentIngestor:
    def __init__(self):
        self.embedder = LiteLLMEmbeddings(model=os.getenv("LITELLM_EMBED_MODEL", "text-embedding-3-large"))

        self.driver = GraphDatabase.driver(
            os.getenv("NEO4J_URI"),
            auth=(os.getenv("NEO4J_USERNAME"), os.getenv("NEO4J_PASSWORD")),
        )

        self.text_splitter = RecursiveCharacterTextSplitter(
            chunk_size=500, chunk_overlap=50, length_function=len, is_separator_regex=False
        )

    def _get_text_preview(self, file_path: str, max_chars: int = 6000) -> str:
        docs = self._load_document(file_path)
        text = " ".join([d.page_content for d in docs])
        return text[:max_chars]

    def _hash_file(self, file_path: str) -> str:
        hasher = hashlib.md5()
        with open(file_path, "rb") as f:
            while chunk := f.read(8192):
                hasher.update(chunk)
        return hasher.hexdigest()

    def _load_document(self, file_path: str):
        ext = os.path.splitext(file_path)[1].lower()
        try:
            if ext == ".pdf":
                return PyMuPDFLoader(file_path).load()
            elif ext == ".docx":
                try:
                    return Docx2txtLoader(file_path).load()
                except Exception as e:
                    print(f"      ⚠️ Docx2txt failed, trying alternative method: {e}")
                    # Fallback: try to extract text using python-docx
                    try:
                        import docx
                        doc = docx.Document(file_path)
                        text = ""
                        for paragraph in doc.paragraphs:
                            text += paragraph.text + "\n"
                        if not text.strip():
                            raise ValueError("No readable text found in document")
                        return [Document(page_content=text, metadata={"source": file_path})]
                    except Exception as e2:
                        print(f"      ⚠️ Alternative docx method also failed: {e2}")
                        # Last resort: return a placeholder document
                        return [Document(page_content="[Document content could not be extracted - file may be corrupted or password protected]", metadata={"source": file_path})]
            elif ext in [".md", ".txt"]:
                try:
                    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                        text = f.read()
                    if not text.strip():
                        return [Document(page_content="[Empty or unreadable file]", metadata={"source": file_path})]
                    return [Document(page_content=text, metadata={"source": file_path})]
                except Exception as e:
                    raise ValueError(f"Failed to load markdown/txt file {file_path}: {e}")
            else:
                raise ValueError(f"Unsupported file type: {ext}")
        except Exception as e:
            print(f"      ❌ Document loading failed: {e}")
            # Return a placeholder document so processing can continue
            return [Document(page_content=f"[Document loading failed: {str(e)}]", metadata={"source": file_path})]

    def process(self, file_path: str):
        """
        Process document and create chunks with embeddings.
        Note: Document node should already exist from main.py processing.
        """
        print(f"      📄 Loading document: {os.path.basename(file_path)}")
        docs = self._load_document(file_path)
        
        print(f"      ✂️ Splitting into chunks...")
        chunks = self.text_splitter.split_documents(docs)
        print(f"      ✅ Created {len(chunks)} chunks")

        file_hash = self._hash_file(file_path)
        file_name = os.path.basename(file_path)

        print(f"      🔗 Linking chunks to existing document node...")
        with self.driver.session() as session:
            # Verify document exists (should be created by main.py)
            doc_check = session.run(
                "MATCH (d:Document {source_id: $source_id}) RETURN d.name AS name",
                source_id=file_hash
            ).single()
            
            if not doc_check:
                print(f"      ⚠️ Document node not found, creating it...")
                session.run(
                    """
                    MERGE (d:Document {source_id: $source_id})
                    SET d.name = $name,
                        d.path = $path,
                        d.processed_at = datetime()
                    """,
                    source_id=file_hash,
                    name=file_name,
                    path=file_path,
                )

            print(f"      🧠 Creating embeddings for chunks...")
            for i, chunk in enumerate(chunks):
                chunk_id = f"{file_hash}_{i}"
                chunk_text = chunk.page_content

                try:
                    embedding = self.embedder.embed_query(chunk_text)
                except Exception as e:
                    print(f"      ❌ Failed to embed chunk {i+1}/{len(chunks)}: {e}")
                    continue

                session.run(
                    """
                    MERGE (c:Chunk {chunk_id: $chunk_id})
                    SET c.text = $text,
                        c.embedding = $embedding,
                        c.chunk_index = $chunk_index
                    MERGE (d:Document {source_id: $source_id})
                    MERGE (d)-[:HAS_CHUNK]->(c)
                    """,
                    chunk_id=chunk_id,
                    text=chunk_text,
                    embedding=embedding,
                    chunk_index=i,
                    source_id=file_hash,
                )

                chunk.metadata["chunk_id"] = chunk_id
                
                if (i + 1) % 10 == 0 or i == len(chunks) - 1:
                    print(f"      📊 Processed {i+1}/{len(chunks)} chunks")

        print(f"      ✅ All chunks linked to document")
        return chunks
